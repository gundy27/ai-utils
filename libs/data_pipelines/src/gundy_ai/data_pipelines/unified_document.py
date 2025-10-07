"""Unified document processor combining basic and advanced functionality."""

from __future__ import annotations

import asyncio
import re
import time
import warnings
from pathlib import Path
from typing import Any

import structlog
import tiktoken
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from PyPDF2 import PdfReader

from .base import BaseProcessor, ProcessingResult, ProcessorConfig
from .document import DocumentMetadata, DocumentType, ProcessedDocument
from .parsers import (
    OCRFallbackParser,
    PDFParser,
    PDFParsingResult,
    PDFPlumberParser,
    PyMuPDFParser,
    PyPDFParser,
)
from .text_cleaning import (
    GeneralTextCleaner,
    OCRTextCleaner,
    PDFTextCleaner,
    TextCleaner,
)
from .audit_integration import ProcessingAuditHook
from .config import get_config

logger = structlog.get_logger(__name__)


class DocumentProcessor(BaseProcessor[str, ProcessedDocument]):
    """Unified document processor with basic and advanced capabilities.

    This processor can operate in two modes:
    1. Basic mode (default): Simple text extraction like the original DocumentProcessor
    2. Advanced mode: Multiple PDF parsers, OCR, text cleaning, audit logging

    The mode is automatically determined based on the parameters provided.
    """

    def __init__(
        self,
        config: ProcessorConfig,
        # Advanced features (when specified, enables advanced mode)
        pdf_parser_priority: list[str] | None = None,
        enable_ocr_fallback: bool = False,
        ocr_confidence_threshold: float = 0.6,
        min_text_extraction_ratio: float = 0.1,
        enable_text_cleaning: bool = False,
        text_cleaning_strategy: str = "auto",
        audit_hook: Any | None = None,
        actor_id: str = "system",
        tenant_id: str = "default",
        # Legacy compatibility
        **kwargs: Any,
    ):
        """Initialize document processor.

        Args:
            config: Processor configuration
            pdf_parser_priority: List of PDF parsers to try in order (enables advanced mode)
            enable_ocr_fallback: Enable OCR for scanned PDFs (enables advanced mode)
            ocr_confidence_threshold: Minimum confidence for OCR text
            min_text_extraction_ratio: Minimum ratio of text to trigger OCR
            enable_text_cleaning: Enable advanced text cleaning (enables advanced mode)
            text_cleaning_strategy: Text cleaning strategy ('auto', 'pdf', 'ocr', 'general')
            audit_hook: Optional audit hook for tracking operations
            actor_id: Actor ID for audit logging
            tenant_id: Tenant ID for audit logging
            **kwargs: Additional arguments for backward compatibility
        """
        super().__init__(config)

        # Determine if we're in advanced mode
        self.advanced_mode = any(
            [
                pdf_parser_priority is not None,
                enable_ocr_fallback,
                enable_text_cleaning,
                # audit_hook is not None  # Temporarily disable audit for refactoring
            ]
        )

        if self.advanced_mode:
            # Advanced mode initialization
            self.pdf_parser_priority = pdf_parser_priority or [
                "pymupdf",
                "pdfplumber",
                "pypdf",
                "ocr_fallback",
            ]
            self.enable_ocr_fallback = enable_ocr_fallback
            self.ocr_confidence_threshold = ocr_confidence_threshold
            self.min_text_extraction_ratio = min_text_extraction_ratio
            self.enable_text_cleaning = enable_text_cleaning
            self.text_cleaning_strategy = text_cleaning_strategy
            self.actor_id = actor_id
            self.tenant_id = tenant_id

            # Initialize audit hook (temporarily disabled for refactoring)
            self.audit_hook = ProcessingAuditHook(None)  # audit_hook

            # Initialize PDF parsers
            self.pdf_parsers: dict[str, PDFParser] = {
                "pypdf": PyPDFParser(),
                "pdfplumber": PDFPlumberParser(),
                "pymupdf": PyMuPDFParser(),
                "ocr_fallback": OCRFallbackParser(
                    confidence_threshold=self.ocr_confidence_threshold
                ),
            }

            # Initialize text cleaners
            self.text_cleaners: dict[str, TextCleaner] = {
                "pdf": PDFTextCleaner(),
                "ocr": OCRTextCleaner(),
                "general": GeneralTextCleaner(),
            }
        else:
            # Basic mode initialization (legacy compatibility)
            self.encoding = tiktoken.get_encoding("cl100k_base")

    def validate_input(self, input_data: str) -> bool:
        """Validate input is a valid file path."""
        if not isinstance(input_data, str):
            return False
        path = Path(input_data)
        return path.exists() and path.is_file()

    async def process(self, file_path: str) -> ProcessingResult[ProcessedDocument]:
        """Process a document file and extract text."""
        if self.advanced_mode:
            return await self._process_advanced(file_path)
        else:
            return await self._process_basic(file_path)

    async def _process_basic(
        self, file_path: str
    ) -> ProcessingResult[ProcessedDocument]:
        """Process document using basic mode (legacy compatibility)."""
        try:
            path = Path(file_path)
            document_type = self._detect_document_type(path)

            # Extract text based on document type
            if document_type == DocumentType.TXT:
                text = await self._extract_txt(path)
            elif document_type == DocumentType.PDF:
                text = await self._extract_pdf_basic(path)
            elif document_type == DocumentType.DOCX:
                text = await self._extract_docx(path)
            elif document_type == DocumentType.HTML:
                text = await self._extract_html(path)
            elif document_type == DocumentType.MARKDOWN:
                text = await self._extract_markdown(path)
            else:
                return ProcessingResult(
                    success=False, error=f"Unsupported document type: {document_type}"
                )

            # Clean and normalize text
            cleaned_text = self._clean_text_basic(text)

            # Create metadata
            metadata = DocumentMetadata(
                filename=path.name,
                document_type=document_type,
                size_bytes=path.stat().st_size,
                word_count=len(cleaned_text.split()),
                custom_metadata={"file_path": str(path), "file_extension": path.suffix},
            )

            # Create processed document
            processed_doc = ProcessedDocument(
                metadata=metadata,
                chunks=[],  # Will be populated by chunking processor
                full_text=cleaned_text,
            )

            return ProcessingResult(
                success=True,
                data=processed_doc,
                metadata={
                    "document_type": document_type.value,
                    "text_length": len(cleaned_text),
                    "word_count": metadata.word_count,
                },
            )

        except Exception as e:
            self.logger.error(
                "document.processing.error", error=str(e), file_path=file_path
            )
            return ProcessingResult(
                success=False, error=f"Failed to process document: {str(e)}"
            )

    async def _process_advanced(
        self, file_path: str
    ) -> ProcessingResult[ProcessedDocument]:
        """Process document using advanced mode with multiple parsers and features."""
        path = Path(file_path)
        document_type = self._detect_document_type(path)

        # Generate event ID for audit trail
        event_id = f"doc_proc_{int(time.time() * 1000)}"

        # Log processing start (temporarily disabled for refactoring)
        # await self.audit_hook.log_processing_start(
        #     file_path=str(path),
        #     actor_id=self.actor_id,
        #     tenant_id=self.tenant_id,
        #     document_type=document_type.value,
        #     processor_type="DocumentProcessor",
        #     event_id=event_id
        # )

        try:
            # Extract text based on document type
            if document_type == DocumentType.PDF:
                text, parsing_metadata = await self._process_pdf_advanced(
                    path, event_id
                )
            elif document_type == DocumentType.TXT:
                text = await self._extract_txt(path)
                parsing_metadata = {"parser_used": "txt_reader"}
            elif document_type == DocumentType.DOCX:
                text = await self._extract_docx(path)
                parsing_metadata = {"parser_used": "docx_reader"}
            elif document_type == DocumentType.HTML:
                text = await self._extract_html(path)
                parsing_metadata = {"parser_used": "html_parser"}
            elif document_type == DocumentType.MARKDOWN:
                text = await self._extract_markdown(path)
                parsing_metadata = {"parser_used": "markdown_reader"}
            else:
                # await self.audit_hook.log_processing_failure(
                #     file_path=str(path),
                #     actor_id=self.actor_id,
                #     tenant_id=self.tenant_id,
                #     error=f"Unsupported document type: {document_type}",
                #     event_id=event_id
                # )
                return ProcessingResult(
                    success=False, error=f"Unsupported document type: {document_type}"
                )

            # Apply text cleaning if enabled
            if self.enable_text_cleaning and text:
                text = await self._apply_text_cleaning(
                    text, document_type, parsing_metadata, event_id
                )

            # Create metadata
            metadata = DocumentMetadata(
                filename=path.name,
                document_type=document_type,
                size_bytes=path.stat().st_size,
                word_count=len(text.split()) if text else 0,
                custom_metadata={
                    "file_path": str(path),
                    "file_extension": path.suffix,
                    "event_id": event_id,
                    **parsing_metadata,
                },
            )

            # Create processed document
            processed_doc = ProcessedDocument(
                metadata=metadata,
                chunks=[],  # Will be populated by chunking processor
                full_text=text,
            )

            # Log successful processing (temporarily disabled for refactoring)
            # await self.audit_hook.log_processing_success(
            #     file_path=str(path),
            #     actor_id=self.actor_id,
            #     tenant_id=self.tenant_id,
            #     text_length=len(text) if text else 0,
            #     word_count=metadata.word_count,
            #     event_id=event_id
            # )

            return ProcessingResult(
                success=True,
                data=processed_doc,
                metadata={
                    "document_type": document_type.value,
                    "text_length": len(text) if text else 0,
                    "word_count": metadata.word_count,
                    "event_id": event_id,
                    **parsing_metadata,
                },
            )

        except Exception as e:
            # await self.audit_hook.log_processing_failure(
            #     file_path=str(path),
            #     actor_id=self.actor_id,
            #     tenant_id=self.tenant_id,
            #     error=str(e),
            #     event_id=event_id
            # )
            self.logger.error(
                "document.processing.error",
                error=str(e),
                file_path=file_path,
                event_id=event_id,
            )
            return ProcessingResult(
                success=False, error=f"Failed to process document: {str(e)}"
            )

    async def _process_pdf_advanced(
        self, path: Path, event_id: str
    ) -> tuple[str, dict[str, Any]]:
        """Process PDF using advanced parsing with multiple parsers and fallback."""
        parsing_attempts = []

        for parser_name in self.pdf_parser_priority:
            if parser_name not in self.pdf_parsers:
                continue

            parser = self.pdf_parsers[parser_name]

            try:
                # Log parsing attempt
                await self.audit_hook.log_parsing_attempt(
                    event_id=event_id,
                    parser_name=parser_name,
                    actor_id=self.actor_id,
                    tenant_id=self.tenant_id,
                )

                result = await parser.parse(path)

                if result.success and result.text:
                    # Check if we got enough text (avoid OCR if not needed)
                    text_ratio = len(result.text.strip()) / max(path.stat().st_size, 1)

                    if (
                        parser_name != "ocr_fallback"
                        and text_ratio < self.min_text_extraction_ratio
                        and self.enable_ocr_fallback
                        and "ocr_fallback" in self.pdf_parser_priority
                    ):
                        # Text extraction ratio too low, continue to OCR
                        parsing_attempts.append(
                            {
                                "parser": parser_name,
                                "success": True,
                                "text_length": len(result.text),
                                "text_ratio": text_ratio,
                                "skipped_reason": "low_text_ratio",
                            }
                        )
                        continue

                    # Success!
                    parsing_attempts.append(
                        {
                            "parser": parser_name,
                            "success": True,
                            "text_length": len(result.text),
                            "confidence": getattr(result, "confidence", None),
                            "pages_processed": getattr(result, "pages_processed", None),
                        }
                    )

                    await self.audit_hook.log_parsing_success(
                        event_id=event_id,
                        parser_name=parser_name,
                        text_length=len(result.text),
                        confidence=getattr(result, "confidence", None),
                        actor_id=self.actor_id,
                        tenant_id=self.tenant_id,
                    )

                    return result.text, {
                        "parser_used": parser_name,
                        "parsing_attempts": parsing_attempts,
                        "confidence": getattr(result, "confidence", None),
                        "pages_processed": getattr(result, "pages_processed", None),
                    }
                else:
                    parsing_attempts.append(
                        {"parser": parser_name, "success": False, "error": result.error}
                    )

            except Exception as e:
                parsing_attempts.append(
                    {"parser": parser_name, "success": False, "error": str(e)}
                )

                await self.audit_hook.log_parsing_failure(
                    event_id=event_id,
                    parser_name=parser_name,
                    error=str(e),
                    actor_id=self.actor_id,
                    tenant_id=self.tenant_id,
                )

        # All parsers failed
        raise Exception(f"All PDF parsers failed. Attempts: {parsing_attempts}")

    async def _apply_text_cleaning(
        self,
        text: str,
        document_type: DocumentType,
        parsing_metadata: dict[str, Any],
        event_id: str,
    ) -> str:
        """Apply text cleaning based on strategy and document type."""
        if not text:
            return text

        # Determine cleaning strategy
        if self.text_cleaning_strategy == "auto":
            # Auto-detect based on document type and parser used
            parser_used = parsing_metadata.get("parser_used", "")
            if document_type == DocumentType.PDF:
                if parser_used == "ocr_fallback":
                    strategy = "ocr"
                else:
                    strategy = "pdf"
            else:
                strategy = "general"
        else:
            strategy = self.text_cleaning_strategy

        if strategy not in self.text_cleaners:
            strategy = "general"  # Fallback

        cleaner = self.text_cleaners[strategy]

        try:
            # Log text cleaning (temporarily disabled for refactoring)
            # await self.audit_hook.log_text_cleaning(
            #     event_id=event_id,
            #     cleaning_strategy=strategy,
            #     original_length=len(text),
            #     actor_id=self.actor_id,
            #     tenant_id=self.tenant_id
            # )

            cleaned_text = await cleaner.clean(text)

            self.logger.info(
                "text_cleaning_applied",
                strategy=strategy,
                original_length=len(text),
                cleaned_length=len(cleaned_text),
                event_id=event_id,
            )

            return cleaned_text

        except Exception as e:
            self.logger.warning(
                "text_cleaning_failed",
                strategy=strategy,
                error=str(e),
                event_id=event_id,
            )
            return text  # Return original text if cleaning fails

    def _detect_document_type(self, path: Path) -> DocumentType:
        """Detect document type from file extension."""
        suffix = path.suffix.lower()

        if suffix == ".txt":
            return DocumentType.TXT
        elif suffix == ".pdf":
            return DocumentType.PDF
        elif suffix in [".docx", ".doc"]:
            return DocumentType.DOCX
        elif suffix in [".html", ".htm"]:
            return DocumentType.HTML
        elif suffix in [".md", ".markdown"]:
            return DocumentType.MARKDOWN
        else:
            raise ValueError(f"Unsupported file extension: {suffix}")

    async def _extract_txt(self, path: Path) -> str:
        """Extract text from plain text file."""
        import asyncio

        return await asyncio.get_event_loop().run_in_executor(
            None, lambda: path.read_text(encoding="utf-8", errors="replace")
        )

    async def _extract_pdf_basic(self, path: Path) -> str:
        """Extract text from PDF file using basic PyPDF2."""
        import asyncio

        def _extract():
            with open(path, "rb") as file:
                reader = PdfReader(file)
                text = ""
                for page in reader.pages:
                    text += page.extract_text() + "\n"
                return text

        return await asyncio.get_event_loop().run_in_executor(None, _extract)

    async def _extract_docx(self, path: Path) -> str:
        """Extract text from DOCX file."""
        import asyncio

        def _extract():
            doc = DocxDocument(path)
            text = ""
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            return text

        return await asyncio.get_event_loop().run_in_executor(None, _extract)

    async def _extract_html(self, path: Path) -> str:
        """Extract text from HTML file."""
        import asyncio

        def _extract():
            content = path.read_text(encoding="utf-8", errors="replace")
            soup = BeautifulSoup(content, "html.parser")
            return soup.get_text()

        return await asyncio.get_event_loop().run_in_executor(None, _extract)

    async def _extract_markdown(self, path: Path) -> str:
        """Extract text from Markdown file."""
        # For now, just read as plain text
        # Could be enhanced with markdown parsing in the future
        return await self._extract_txt(path)

    def _clean_text_basic(self, text: str) -> str:
        """Basic text cleaning for legacy compatibility."""
        if not text:
            return ""

        # Remove excessive whitespace
        text = re.sub(r"\s+", " ", text)

        # Remove leading/trailing whitespace
        text = text.strip()

        return text


# Backward compatibility alias with deprecation warning
class EnhancedDocumentProcessor(DocumentProcessor):
    """Deprecated: Use DocumentProcessor with advanced features enabled instead.

    This class is maintained for backward compatibility but will be removed in a future version.
    Use DocumentProcessor with the appropriate parameters to enable advanced features.
    """

    def __init__(self, *args, **kwargs):
        warnings.warn(
            "EnhancedDocumentProcessor is deprecated. Use DocumentProcessor with advanced "
            "features enabled instead. This class will be removed in a future version.",
            DeprecationWarning,
            stacklevel=2,
        )

        # Force advanced mode by setting default advanced parameters
        kwargs.setdefault(
            "pdf_parser_priority", ["pymupdf", "pdfplumber", "pypdf", "ocr_fallback"]
        )
        kwargs.setdefault("enable_ocr_fallback", True)
        kwargs.setdefault("enable_text_cleaning", True)

        super().__init__(*args, **kwargs)
