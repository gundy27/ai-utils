"""Tests for DOCX parser."""

import tempfile
from pathlib import Path

import pytest

from gundy_ai.extractors import DOCXParser, ParsedChunk


@pytest.fixture
def docx_parser():
    """Create DOCX parser instance."""
    return DOCXParser()


def test_docx_parser_manifest(docx_parser):
    """Test DOCX parser manifest."""
    manifest = docx_parser.manifest

    assert manifest.name == "docx"
    assert manifest.version == "1.0.0"
    assert ".docx" in manifest.supported_types
    assert manifest.schema_version == "1.0"
    assert "python-docx>=1.0" in manifest.dependencies


def test_docx_parser_health_check(docx_parser):
    """Test DOCX parser health check."""
    # Should be True since python-docx is installed
    assert docx_parser.health_check() is True


def test_parse_file_not_found(docx_parser):
    """Test parsing non-existent file raises FileNotFoundError."""
    with pytest.raises(FileNotFoundError, match="File not found"):
        docx_parser.parse("/path/to/nonexistent/file.docx")


def test_parse_simple_docx(docx_parser):
    """Test parsing a simple DOCX file."""
    from docx import Document

    # Create a simple DOCX with paragraphs
    doc = Document()
    doc.add_paragraph("First paragraph")
    doc.add_paragraph("Second paragraph")
    doc.add_paragraph("")  # Empty paragraph should be skipped
    doc.add_paragraph("Third paragraph")

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        doc.save(f.name)
        temp_path = f.name

    try:
        chunks = docx_parser.parse(temp_path)

        # Should have 3 chunks (empty paragraph skipped)
        assert len(chunks) == 3
        assert all(isinstance(chunk, ParsedChunk) for chunk in chunks)

        # Check content
        assert chunks[0].text == "First paragraph"
        assert chunks[1].text == "Second paragraph"
        assert chunks[2].text == "Third paragraph"

        # Check metadata
        assert all(chunk.metadata["parser"] == "docx" for chunk in chunks)
        assert all(chunk.metadata["type"] == "paragraph" for chunk in chunks)
        assert chunks[0].metadata["index"] == 0
        assert chunks[1].metadata["index"] == 1
    finally:
        Path(temp_path).unlink(missing_ok=True)


def test_parse_docx_with_table(docx_parser):
    """Test parsing DOCX with table."""
    from docx import Document

    doc = Document()
    doc.add_paragraph("Introduction")

    # Add a table
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Age"
    table.cell(1, 0).text = "Alice"
    table.cell(1, 1).text = "30"
    table.cell(2, 0).text = "Bob"
    table.cell(2, 1).text = "25"

    doc.add_paragraph("Conclusion")

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        doc.save(f.name)
        temp_path = f.name

    try:
        chunks = docx_parser.parse(temp_path)

        # Should have paragraphs + table rows
        assert len(chunks) >= 2

        # Check paragraph chunks
        para_chunks = [c for c in chunks if c.metadata.get("type") == "paragraph"]
        assert len(para_chunks) == 2
        assert "Introduction" in para_chunks[0].text
        assert "Conclusion" in para_chunks[1].text

        # Check table chunks
        table_chunks = [c for c in chunks if c.metadata.get("type") == "table_row"]
        assert len(table_chunks) == 3

        # Header row should contain Name and Age
        assert "Name" in table_chunks[0].text
        assert "Age" in table_chunks[0].text

        # Combine all table text to verify all data is present
        all_table_text = " ".join(c.text for c in table_chunks)
        assert "Name" in all_table_text
        assert "Age" in all_table_text
        assert "Alice" in all_table_text
        assert "30" in all_table_text
        assert "Bob" in all_table_text
        assert "25" in all_table_text
    finally:
        Path(temp_path).unlink(missing_ok=True)


def test_parse_docx_with_merged_cells(docx_parser):
    """Test that merged cells don't create duplicates."""
    from docx import Document

    doc = Document()
    table = doc.add_table(rows=2, cols=2)

    # Merge cells in first row
    cell_a = table.cell(0, 0)
    cell_b = table.cell(0, 1)
    cell_a.merge(cell_b)
    cell_a.text = "Merged Header"

    table.cell(1, 0).text = "Data 1"
    table.cell(1, 1).text = "Data 2"

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        doc.save(f.name)
        temp_path = f.name

    try:
        chunks = docx_parser.parse(temp_path)

        table_chunks = [c for c in chunks if c.metadata.get("type") == "table_row"]

        # First row should only have "Merged Header" once
        assert "Merged Header" in table_chunks[0].text
        # Count occurrences
        assert table_chunks[0].text.count("Merged Header") == 1

        # Second row should have both cells
        assert "Data 1" in table_chunks[1].text
        assert "Data 2" in table_chunks[1].text
    finally:
        Path(temp_path).unlink(missing_ok=True)


def test_parse_empty_docx(docx_parser):
    """Test parsing empty DOCX."""
    from docx import Document

    doc = Document()

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        doc.save(f.name)
        temp_path = f.name

    try:
        chunks = docx_parser.parse(temp_path)

        # Empty document should return empty list
        assert chunks == []
    finally:
        Path(temp_path).unlink(missing_ok=True)


def test_parse_invalid_docx(docx_parser):
    """Test parsing invalid DOCX file."""
    with tempfile.NamedTemporaryFile(mode="w", suffix=".docx", delete=False) as f:
        f.write("This is not a DOCX file")
        temp_path = f.name

    try:
        with pytest.raises(RuntimeError, match="Failed to parse DOCX"):
            docx_parser.parse(temp_path)
    finally:
        Path(temp_path).unlink(missing_ok=True)


def test_parsed_chunk_metadata_structure(docx_parser):
    """Test that chunks have expected metadata."""
    from docx import Document

    doc = Document()
    doc.add_paragraph("Test content")

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        doc.save(f.name)
        temp_path = f.name

    try:
        chunks = docx_parser.parse(temp_path)

        metadata = chunks[0].metadata
        assert "parser" in metadata
        assert "parser_version" in metadata
        assert "type" in metadata
        assert "index" in metadata
        assert "file_name" in metadata
        assert "file_extension" in metadata

        assert metadata["parser"] == "docx"
        assert metadata["parser_version"] == "1.0.0"
        assert metadata["type"] == "paragraph"
        assert metadata["file_extension"] == ".docx"
    finally:
        Path(temp_path).unlink(missing_ok=True)


def test_span_offsets_sequential(docx_parser):
    """Test that span offsets are sequential."""
    from docx import Document

    doc = Document()
    doc.add_paragraph("First")
    doc.add_paragraph("Second")
    doc.add_paragraph("Third")

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        doc.save(f.name)
        temp_path = f.name

    try:
        chunks = docx_parser.parse(temp_path)

        # Spans should be sequential
        assert chunks[0].span_start == 0
        assert chunks[0].span_end == len("First")
        assert chunks[1].span_start == chunks[0].span_end
        assert chunks[1].span_end == chunks[1].span_start + len("Second")
    finally:
        Path(temp_path).unlink(missing_ok=True)
