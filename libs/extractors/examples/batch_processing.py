"""Batch processing example for multiple documents."""

from pathlib import Path
import tempfile

from gundy_ai.extractors import DOCXParser, PDFParser, ParserRegistry, TXTParser


def main():
    """Demonstrate batch processing of multiple documents."""

    # Create registry with all parsers
    registry = ParserRegistry()
    registry.register(TXTParser())
    registry.register(PDFParser())
    registry.register(DOCXParser())

    print("Registered Parsers:")
    for manifest in registry.list_parsers():
        print(
            f"  {manifest.name} v{manifest.version}: {', '.join(manifest.supported_types)}"
        )

    # Create sample files
    print("\nCreating sample files...")

    # TXT file
    txt_file = tempfile.NamedTemporaryFile(mode="w", suffix=".txt", delete=False)
    txt_file.write("Sample text document\nWith multiple lines")
    txt_file.close()

    # PDF file
    from pypdf import PdfWriter

    pdf_file = tempfile.NamedTemporaryFile(suffix=".pdf", delete=False)
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.write(pdf_file)
    pdf_file.close()

    # DOCX file
    from docx import Document

    docx_file = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    doc = Document()
    doc.add_paragraph("Sample Word document")
    doc.add_paragraph("With formatted text")
    doc.save(docx_file.name)
    docx_file.close()

    # Process all files
    files = [txt_file.name, pdf_file.name, docx_file.name]

    print(f"\nProcessing {len(files)} files...")

    total_chunks = 0
    total_characters = 0

    for file_path in files:
        path = Path(file_path)
        extension = path.suffix

        # Get appropriate parser
        parser = registry.get_parser(extension)
        if parser is None:
            print(f"  ⚠ No parser for {extension}: {path.name}")
            continue

        try:
            # Parse file
            chunks = parser.parse(file_path)
            chars = sum(len(chunk.text) for chunk in chunks)

            total_chunks += len(chunks)
            total_characters += chars

            print(f"  ✓ {path.name} ({extension}): {len(chunks)} chunks, {chars} chars")

        except Exception as e:
            print(f"  ✗ {path.name}: {str(e)}")

    print("\nBatch Processing Complete:")
    print(f"  Total files: {len(files)}")
    print(f"  Total chunks: {total_chunks}")
    print(f"  Total characters: {total_characters}")

    # Cleanup
    for file_path in files:
        Path(file_path).unlink(missing_ok=True)


if __name__ == "__main__":
    main()
