"""Example demonstrating all available parsers."""

from pathlib import Path
import tempfile

from gundy_ai.extractors import DOCXParser, PDFParser, ParserRegistry, TXTParser


def main():
    """Demonstrate all three parsers (TXT, PDF, DOCX)."""

    print("=== Gundy AI Extractors - All Parsers Demo ===\n")

    # Create registry with all parsers
    registry = ParserRegistry()
    registry.register(TXTParser())
    registry.register(PDFParser())
    registry.register(DOCXParser())

    print("Registered Parsers:")
    for manifest in registry.list_parsers():
        print(
            f"  {manifest.name:6} v{manifest.version:6} - {', '.join(manifest.supported_types)}"
        )

    print(
        f"\nSupported Extensions: {', '.join(sorted(registry.get_supported_extensions()))}"
    )

    # Health check
    print("\nHealth Check:")
    health_results = registry.health_check_all()
    for name, is_healthy in health_results.items():
        status = "✓ Healthy" if is_healthy else "✗ Unhealthy"
        print(f"  {name}: {status}")

    print("\n" + "=" * 60)

    # 1. TXT Parser
    print("\n1. TEXT PARSER (.txt, .md, .csv)")
    print("-" * 60)

    txt_file = tempfile.NamedTemporaryFile(mode="w", suffix=".txt", delete=False)
    txt_file.write("This is a text document.\n")
    txt_file.write("It has multiple lines of content.\n")
    txt_file.write("The TXT parser reads it all as one chunk.")
    txt_file.close()

    try:
        parser = registry.get_parser(".txt")
        chunks = parser.parse(txt_file.name)

        print(f"File: {Path(txt_file.name).name}")
        print(f"Chunks: {len(chunks)}")
        print(f"Characters: {len(chunks[0].text)}")
        print(f"Preview: {chunks[0].text[:80]}...")
    finally:
        Path(txt_file.name).unlink(missing_ok=True)

    # 2. PDF Parser
    print("\n2. PDF PARSER (.pdf)")
    print("-" * 60)

    from pypdf import PdfWriter

    pdf_file = tempfile.NamedTemporaryFile(suffix=".pdf", delete=False)
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.add_blank_page(width=200, height=200)
    writer.write(pdf_file)
    pdf_file.close()

    try:
        parser = registry.get_parser(".pdf")
        chunks = parser.parse(pdf_file.name)

        print(f"File: {Path(pdf_file.name).name}")
        print(f"Chunks: {len(chunks)} (one per page)")
        for i, chunk in enumerate(chunks, 1):
            print(
                f"  Page {i}: {chunk.metadata['page']}/{chunk.metadata['total_pages']}"
            )
    finally:
        Path(pdf_file.name).unlink(missing_ok=True)

    # 3. DOCX Parser
    print("\n3. DOCX PARSER (.docx)")
    print("-" * 60)

    from docx import Document

    docx_file = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    doc = Document()
    doc.add_paragraph("Introduction paragraph")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Header 1"
    table.cell(0, 1).text = "Header 2"
    table.cell(1, 0).text = "Data 1"
    table.cell(1, 1).text = "Data 2"

    doc.add_paragraph("Conclusion paragraph")
    doc.save(docx_file.name)
    docx_file.close()

    try:
        parser = registry.get_parser(".docx")
        chunks = parser.parse(docx_file.name)

        print(f"File: {Path(docx_file.name).name}")
        print(f"Chunks: {len(chunks)}")

        para_count = sum(1 for c in chunks if c.metadata.get("type") == "paragraph")
        table_count = sum(1 for c in chunks if c.metadata.get("type") == "table_row")

        print(f"  Paragraphs: {para_count}")
        print(f"  Table rows: {table_count}")

        print("\nChunk breakdown:")
        for i, chunk in enumerate(chunks, 1):
            chunk_type = chunk.metadata.get("type", "unknown")
            preview = chunk.text[:40] + "..." if len(chunk.text) > 40 else chunk.text
            print(f"  {i}. [{chunk_type}] {preview}")
    finally:
        Path(docx_file.name).unlink(missing_ok=True)

    print("\n" + "=" * 60)
    print("\n✓ All parsers working!")
    print("\nNext steps:")
    print("  - Integrate with chunking system")
    print("  - Add embedding generation")
    print("  - Store in vector database")


if __name__ == "__main__":
    main()
